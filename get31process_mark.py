# -*- coding: utf-8 -*-
import glob
import os
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

# ===================== 配置区（只改这里） =====================
# 市场选择：US / A / HK（改这一个变量即可切换市场，或通过环境变量 MARKET 传入）
MARKET = os.environ.get("MARKET", "US")

# 基础路径（自动根据 MARKET 生成输入/输出目录）
BASE_DIR = r"/home/ubuntu/alibaba_agent/downlaods_futu_data/data_all_market"
INPUT_DIR = os.path.join(BASE_DIR, MARKET, "demo", "计算区间", "双区间相同再算幅度", "标黄")           # 输入目录
OUTPUT_DIR = os.path.join(BASE_DIR, MARKET, "demo", "计算区间", "双区间相同再算幅度", "标黄", "已标注")  # 输出目录
# ============================================================

SHEET_NAME = '双区间再分表格'
COL_A = 1
COL_G = 7
COL_H = 8

RED_FILL = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')
YELLOW_FILL = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')

MARKED_SUFFIX = '_已标注.xlsx'
RED_SUFFIX = '_标红色.txt'
YELLOW_SUFFIX = '_标黄色.txt'

def is_yellow_fill(cell):
    fill = cell.fill
    if not fill or fill.fill_type is None or fill.fill_type == 'none':
        return False
    fg = fill.fgColor
    if fg is None:
        return False
    if fg.type == 'rgb':
        rgb_str = str(fg.rgb).upper()
        return rgb_str[-6:] == 'FFFF00'
    return False

def strip_prefix(code):
    if not code:
        return ''
    s = str(code).strip()
    for prefix in ('US.', 'HK.', 'SZ.', 'SH.'):
        if s.upper().startswith(prefix):
            return s[len(prefix):]
    return s

def get_output_paths(xlsx_path, output_dir):
    src = Path(xlsx_path)
    stem = src.stem
    out_excel = os.path.join(output_dir, f'{stem}_已标注.xlsx')
    out_red = os.path.join(output_dir, f'{stem}_标红色.txt')
    out_yellow = os.path.join(output_dir, f'{stem}_标黄色.txt')
    return out_excel, out_red, out_yellow

def write_codes(txt_path, codes):
    with open(txt_path, 'w', encoding='utf-8') as f:
        for code in codes:
            f.write(code + '\n')

def process_one_file(xlsx_path, output_dir):
    print(f"\n{'=' * 60}")
    print(f"Processing: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    if SHEET_NAME not in wb.sheetnames:
        print(f"  [SKIP] Sheet '{SHEET_NAME}' not found. Available: {wb.sheetnames}")
        wb.close()
        return 0, 0

    sheet = wb[SHEET_NAME]
    max_row = sheet.max_row
    max_col = sheet.max_column
    print(f"  Sheet '{SHEET_NAME}': rows={max_row}, cols={max_col}")

    red_codes = []
    yellow_codes = []

    for row_idx in range(2, max_row + 1):
        a_cell = sheet.cell(row=row_idx, column=COL_A)
        a_val = a_cell.value

        g_yellow = is_yellow_fill(sheet.cell(row=row_idx, column=COL_G))
        h_yellow = is_yellow_fill(sheet.cell(row=row_idx, column=COL_H))

        if g_yellow and h_yellow:
            a_cell.fill = RED_FILL
            if a_val:
                red_codes.append(strip_prefix(a_val))
            continue

        found_later_pair = False
        for col_idx in range(COL_H, max_col):
            c1 = is_yellow_fill(sheet.cell(row=row_idx, column=col_idx))
            c2 = is_yellow_fill(sheet.cell(row=row_idx, column=col_idx + 1))
            if c1 and c2:
                found_later_pair = True
                break

        if found_later_pair:
            a_cell.fill = YELLOW_FILL
            if a_val:
                yellow_codes.append(strip_prefix(a_val))

    red_clean = list(dict.fromkeys([c for c in red_codes if c]))
    yellow_clean = list(dict.fromkeys([c for c in yellow_codes if c]))

    out_excel, red_txt, yellow_txt = get_output_paths(xlsx_path, output_dir)
    print(f"  Saving Excel to: {out_excel}")
    print(f"  Saving red list to: {red_txt}")
    print(f"  Saving yellow list to: {yellow_txt}")

    wb.save(out_excel)
    wb.close()

    write_codes(red_txt, red_clean)
    write_codes(yellow_txt, yellow_clean)

    print(f"  红色 {len(red_clean)} 条，黄色 {len(yellow_clean)} 条")
    return len(red_clean), len(yellow_clean)

def get_target_files(input_path):
    if not os.path.exists(input_path):
        print(f"INPUT_DIR not found: {input_path}")
        return []

    if os.path.isfile(input_path):
        if not input_path.lower().endswith('.xlsx'):
            print(f"INPUT_PATH is a file but not xlsx: {input_path}")
            return []
        return [input_path]

    if os.path.isdir(input_path):
        all_xlsx = sorted(glob.glob(os.path.join(input_path, '*.xlsx')))
        target_files = []
        for p in all_xlsx:
            name = os.path.basename(p)
            if name.startswith('~$'):
                continue
            if name.endswith(MARKED_SUFFIX):
                continue
            target_files.append(p)
        return target_files

    return []

def main():
    # 如果输出目录不存在则创建
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    print(f"📂 输入目录: {INPUT_DIR}")
    print(f"📂 输出目录: {OUTPUT_DIR}")
    
    target_files = get_target_files(INPUT_DIR)

    if not target_files:
        print("No xlsx files to process.")
        return

    print(f"Found {len(target_files)} xlsx file(s) to process.")

    total_red = 0
    total_yellow = 0
    for xlsx_path in target_files:
        r, y = process_one_file(xlsx_path, OUTPUT_DIR)
        total_red += r
        total_yellow += y

    print(f"\n{'=' * 60}")
    print(f"ALL DONE. Files processed: {len(target_files)}")
    print(f"  红色总计: {total_red} 条")
    print(f"  黄色总计: {total_yellow} 条")

if __name__ == '__main__':
    main()
