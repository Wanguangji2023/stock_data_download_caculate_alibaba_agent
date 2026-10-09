# -*- coding: utf-8 -*-
"""
修复版 utilnew_yun.py：读取 CSV / XLSX 获取股票代码列表
======================================================
修复内容（对照原版）：
  1. 【关键】encoding="atin-1" 拼写错误 -> 改为自动检测编码
     （按 检测编码 -> utf-8-sig -> gb18030 -> latin-1 逐个尝试）
  2. 列名容错：代码列 / 名称列支持中文与英文多种写法，找不到用位置列
  3. 统一过滤空行 / NaN / 末尾 ".0"（A股、港股、美股）
  4. readCSV_US 的“名称”列改为可选（找不到不再 KeyError）
  5. ST 过滤更精确：名称含 “ *ST” / 以“ST ”开头 / 代码带“ST-”前缀才过滤，
     不再误伤 STMicroelectronics 这类名称含 ST 的正常股
  6. A股区分 SH(6) / SZ(0,3) / BJ(4,8)，不再把所有非6开头都当 SZ
"""
import pandas as pd
from openpyxl import load_workbook

pd.set_option("display.max_rows", None)
pd.set_option("display.max_columns", None)

# 代码列 / 名称列的候选列名（按优先级）
_CODE_COLS = ["代码", "证券代码", "股票代码", "code", "Code", "CODE"]
_NAME_COLS = ["名称", "股票名称", "证券简称", "name", "Name", "NAME"]


# ============================================================
# 内部工具
# ============================================================
def _detect_encoding(path):
    """检测文件编码，优先 chardet，失败默认 utf-8"""
    try:
        import chardet
        with open(path, "rb") as f:
            raw = f.read(1024 * 1024)
        return chardet.detect(raw).get("encoding") or "utf-8"
    except Exception:
        return "utf-8"


def _read_df(csvfilepath):
    """按 检测编码 -> utf-8-sig -> gb18030 -> latin-1 顺序尝试读取，dtype=str 保留前导零"""
    enc = _detect_encoding(csvfilepath)
    for e in dict.fromkeys([enc, "utf-8-sig", "gb18030", "latin-1"]):
        try:
            return pd.read_csv(csvfilepath, encoding=e, dtype=str)
        except Exception:
            continue
    raise ValueError(f"无法解码 CSV：{csvfilepath}")


def _find_col(df, candidates):
    """按候选列名查找列（去空白、忽略大小写），找不到返回 None"""
    exact = {str(c) for c in df.columns}
    for c in candidates:
        if c in exact:
            return c
    lowered = {str(c).strip().lower(): c for c in df.columns}
    for c in candidates:
        if c.strip().lower() in lowered:
            return lowered[c.strip().lower()]
    return None


def _is_st(name, code):
    """判断是否退市风险 ST 股票（更精确，避免误伤名称含 ST 字母的个股）"""
    n = (name or "").upper()
    c = (code or "").upper()
    return " *ST" in n or n.startswith("ST ") or c.startswith("ST-") or c.startswith("ST ")


def _clean_codes(df, code_col, name_col=None, filter_st=False):
    """提取干净代码列表：过滤空值 / NaN / 末尾 .0；可选过滤 ST"""
    out = []
    for idx, row in df.iterrows():
        raw = str(row.get(code_col, "")).strip()
        if not raw or raw.lower() == "nan":
            continue
        if raw.endswith(".0"):
            raw = raw[:-2]
        raw = raw.strip()
        if not raw or raw.lower() == "nan":
            continue
        if filter_st and name_col is not None:
            name = str(row.get(name_col, "")).strip()
            if _is_st(name, raw):
                print(f"跳过ST_{name}\n")
                continue
        out.append(raw)
    return out


# ============================================================
# A股：代码按 SH(6) / SZ(0,3) / BJ(4,8) 加前缀
# ============================================================
def readCSV_A(csvfilepath):
    df = _read_df(csvfilepath)
    code_col = _find_col(df, _CODE_COLS) or df.columns[0]
    codes = _clean_codes(df, code_col)
    out = []
    for t in codes:
        t = t.zfill(6)
        if t.startswith("6"):
            out.append("SH." + t)
        elif t.startswith("4") or t.startswith("8"):
            out.append("BJ." + t)
        else:
            out.append("SZ." + t)
    print(f"pd所有数据行数：{len(df)}，有效代码：{len(out)}")
    return out


# ============================================================
# 港股：代码补零成 5 位
# ============================================================
def readCSV_HK(csvfilepath):
    df = _read_df(csvfilepath)
    code_col = _find_col(df, _CODE_COLS) or df.columns[0]
    codes = _clean_codes(df, code_col)
    out = ["HK." + t.zfill(5) for t in codes]
    print(f"pd所有数据行数：{len(df)}，有效代码：{len(out)}")
    return out


# ============================================================
# 美股：不过滤补零，可选 ST 过滤（名称列缺失时自动跳过 ST 判断）
# ============================================================
def readCSV_US(csvfilepath):
    df = _read_df(csvfilepath)
    code_col = _find_col(df, _CODE_COLS) or df.columns[0]
    name_col = _find_col(df, _NAME_COLS)
    codes = _clean_codes(df, code_col, name_col, filter_st=True)
    out = ["US." + t for t in codes]
    print(f"pd所有数据行数：{len(df)}，有效代码：{len(out)}")
    return out


# ============================================================
# XLSX：读取第一个工作表指定列所有非空数据（去掉表头行）
# ============================================================
def get_excel_first_sheet_column(file_path: str, target_col: int) -> list:
    """
    读取xlsx第一个工作表指定列所有非空数据，返回列表
    :param file_path: xlsx文件完整路径
    :param target_col: 目标列数字，A=1, B=2, H=8...
    :return: 该列所有有效数据list（不含表头）
    """
    data_list = []
    wb = load_workbook(file_path, data_only=True)
    ws = wb.worksheets[0]
    for row in range(1, ws.max_row + 1):
        cell_val = ws.cell(row=row, column=target_col).value
        if cell_val is not None:
            data_list.append(cell_val)
    wb.close()
    return list(data_list[1:])
